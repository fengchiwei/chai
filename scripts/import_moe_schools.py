"""Build data/schools_moe.json from the Ministry of Education school directories (教育部統計處「各級學校名錄」).

Run by the team when a new academic year is published; the app only reads the JSON. Needs openpyxl
(not a runtime dependency):

    python scripts/import_moe_schools.py            # download the 115 學年 files and rebuild
    python scripts/import_moe_schools.py --year 116
    python scripts/import_moe_schools.py --from-dir path/to/xlsx   # use files already downloaded

The directories list each school's code, name, public/private, county, address, main phone and website.
They do not list counseling-office extensions or emails; those stay in data/school_counseling.json, which
only holds contacts the team has verified.
"""
import argparse
import io
import json
import os
import re
import sys
from datetime import date

import openpyxl
import requests

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "data", "schools_moe.json")
PAGE = "https://depart.moe.edu.tw/ED4500/News.aspx?n=63F5AB3D02A8BBAC&sms=1FF9979D10DBF9F3"
FILES = {  # file stem -> school level
    "e1_new": "國小", "j1_new": "國中", "high": "高中職", "sp1_new": "特殊教育學校", "u1_new": "大專校院",
}
HEADERS = {"code": "代碼", "name": "學校名稱", "public": "公/私立", "county": "縣市名稱", "address": "地址",
           "phone": "電話", "website": "網址"}


def rows(book):
    sheet = book.active
    header, out = None, []
    for row in sheet.iter_rows(values_only=True):
        cells = [str(c).strip() if c is not None else "" for c in row]
        if header is None:
            if "代碼" in cells and "學校名稱" in cells:
                header = {key: cells.index(label) for key, label in HEADERS.items()}
            continue
        if not any(cells):
            continue
        record = {key: cells[i] if i < len(cells) else "" for key, i in header.items()}
        county_at = header["county"]
        if re.fullmatch(r"\[\d+\]", record["county"]) and county_at + 1 < len(cells):
            record["county"] = cells[county_at + 1]  # 大專校院: the code and the county name are in two columns
        out.append(record)
    return out


def clean(record, level):
    county = re.sub(r"^\[\d+\]", "", record["county"]).strip()
    address = re.sub(r"^\[\d+\]", "", record["address"]).strip()
    website = record["website"].strip()
    if website and not re.match(r"https?://", website):
        website = "http://" + website
    return {"code": record["code"], "name": record["name"], "level": level, "public": record["public"],
            "county": county, "address": address, "phone": record["phone"], "website": website}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--year", default="115")
    parser.add_argument("--from-dir")
    args = parser.parse_args()
    schools, sources = [], []
    for stem, level in FILES.items():
        url = f"https://stats.moe.gov.tw/files/school/{args.year}/{stem}.xlsx"
        if args.from_dir:
            book = openpyxl.load_workbook(os.path.join(args.from_dir, stem + ".xlsx"), read_only=True)
        else:
            response = requests.get(url, timeout=60)
            response.raise_for_status()
            book = openpyxl.load_workbook(io.BytesIO(response.content), read_only=True)
        found = [clean(r, level) for r in rows(book) if r["code"] and r["name"]]
        if not found:
            sys.exit(f"{stem}: no rows found; the file layout may have changed")
        schools += found
        sources.append({"level": level, "url": url, "count": len(found)})
    data = {
        "_說明": "教育部統計處各級學校名錄（學校總機、地址、官網），由 scripts/import_moe_schools.py 產生，請勿手動修改。"
                 "名錄沒有輔導室分機或信箱；核對過的輔導室信箱放在 data/school_counseling.json。",
        "academic_year": args.year, "retrieved_on": date.today().isoformat(), "source_page": PAGE,
        "sources": sources, "schools": schools,
    }
    with open(OUT, "w", encoding="utf-8", newline="\n") as f:
        json.dump(data, f, ensure_ascii=False, separators=(",", ":"))
        f.write("\n")
    print(f"{len(schools)} schools written to {os.path.relpath(OUT, ROOT)}")
    for s in sources:
        print(f"  {s['level']}: {s['count']}")


if __name__ == "__main__":
    main()
